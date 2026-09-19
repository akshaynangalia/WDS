"""
Run Report tab: which run produced this workbook, on what inputs, how it went,
and whether it passed its health checks -- readable without opening a log.

Two columns (item | value), grouped under purple section bars in the same
style as the other tabs. It is appended after the four approved sheets so none
of their positions move, and it is only written when the run supplies a
RunReport (see output/excel_writer.py).

If a health check failed, a red banner sits at the top of the sheet: the run
still produced its output, but the warning is loud (client decision).

Contract:
    consumes: an openpyxl Workbook, a RunReport
    produces: a new "Run Report" worksheet in that workbook
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.workbook import Workbook

from engine.run_report import RunReport
from output import styling

SHEET_NAME = "Run Report"
_MAX_FALLBACKS_SHOWN = 20
_MAX_CAPACITY_SHOWN = 10

_FAIL_FILL = PatternFill(start_color="FFF8D7DA", end_color="FFF8D7DA", fill_type="solid")
_PASS_FILL = PatternFill(start_color="FFDFF0D8", end_color="FFDFF0D8", fill_type="solid")
_BANNER_FILL = PatternFill(start_color="FFC0392B", end_color="FFC0392B", fill_type="solid")
_WRAP = Alignment(wrap_text=True, vertical="top")

_STATUS_TEXT = {
    "success": "Success",
    "degraded": "Degraded - fallback defaults were applied (see the Assumption Applied tab)",
}


def _section(ws, row: int, title: str) -> int:
    for col in (1, 2):
        cell = ws.cell(row=row, column=col, value=title if col == 1 else None)
        cell.fill = styling.HEADER_FILL
        cell.font = styling.HEADER_FONT
        cell.border = styling.HEADER_BORDER
    return row + 1


def _item(ws, row: int, label: str, value, fill=None) -> int:
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)
    cell = ws.cell(row=row, column=2, value=value)
    cell.alignment = _WRAP
    if fill is not None:
        cell.fill = fill
    return row + 1


def write(workbook: Workbook, report: RunReport) -> None:
    ws = workbook.create_sheet(SHEET_NAME)  # appended last: existing sheet positions do not move
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110

    ws.cell(row=1, column=1, value="Run Report").font = Font(bold=True, size=14)
    row = 3

    failed = [] if report.health is None else [c for c in report.health.checks if not c.ok]
    if failed:
        summary = "; ".join(f"{c.name} ({len(c.failures)} row(s))" for c in failed)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        banner = ws.cell(row=row, column=1,
                         value=f"INTEGRITY WARNING - health check failed: {summary}. "
                               f"The plan was still produced; review before use. Log reference: {report.reference}")
        banner.fill = _BANNER_FILL
        banner.font = Font(bold=True, color="FFFFFFFF")
        banner.alignment = _WRAP
        ws.row_dimensions[row].height = 34
        row += 2

    row = _section(ws, row, "Run summary")
    row = _item(ws, row, "Reference", report.reference)
    row = _item(ws, row, "Generated (UTC)", report.generated_utc)
    row = _item(ws, row, "Tool version", report.version)
    row = _item(ws, row, "Status", _STATUS_TEXT.get(report.status, report.status))
    row = _item(ws, row, "Periods", report.periods)
    row = _item(ws, row, "Lines", report.lines)
    if report.min_dos_entered is not None:
        row = _item(ws, row, "Minimum DOS override entered",
                    f"{report.min_dos_entered:g} days - recorded only; this version of the engine does not apply it")
    row = _item(ws, row, "Elapsed until this report was written", f"{report.elapsed_seconds:.1f} s")
    row += 1

    row = _section(ws, row, "Inputs (identified by fingerprint; the files themselves are not stored)")
    for label, fingerprint in report.inputs:
        row = _item(ws, row, label, fingerprint)
    row += 1

    row = _section(ws, row, "Health checks")
    if report.health is None:
        row = _item(ws, row, "All checks", "NOT RUN - the checks could not run; see the log", _FAIL_FILL)
    else:
        for check in report.health.checks:
            row = _item(ws, row, check.name, f"{'PASS' if check.ok else 'FAIL'} - {check.detail}",
                        _PASS_FILL if check.ok else _FAIL_FILL)
    row += 1

    row = _section(ws, row, "Stage timings")
    for stage in report.stages:
        runs = "1 run" if stage.runs == 1 else f"{stage.runs} runs"
        row = _item(ws, row, stage.name, f"{runs}, {stage.seconds:.2f} s total")
    row += 1

    row = _section(ws, row, "Data quality")
    row = _item(ws, row, "Run-level fallbacks applied", len(report.fallback_messages))
    for message in report.fallback_messages[:_MAX_FALLBACKS_SHOWN]:
        row = _item(ws, row, "", message)
    if len(report.fallback_messages) > _MAX_FALLBACKS_SHOWN:
        row = _item(ws, row, "", f"... and {len(report.fallback_messages) - _MAX_FALLBACKS_SHOWN} more")
    row = _item(ws, row, "Capacity fallbacks applied", len(report.capacity_messages))
    for message in report.capacity_messages[:_MAX_CAPACITY_SHOWN]:
        row = _item(ws, row, "", message)
    if len(report.capacity_messages) > _MAX_CAPACITY_SHOWN:
        row = _item(ws, row, "",
                    f"... and {len(report.capacity_messages) - _MAX_CAPACITY_SHOWN} more "
                    f"(the full list is on the Assumption Applied tab)")
