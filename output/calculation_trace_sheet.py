"""
Calculation Trace tab: for every Link Code in every period, the inputs that
drove the number, which rule fired, and a one-line explanation -- so a planner
can spot a bad input and a developer can spot a logic problem.

Long format (one row per Link Code per period), sorted Plant / Line / Linkcode /
period, with a frozen header and a filter, so any Link Code can be picked out.

The explanation is BUILT HERE, at output time, from facts the engine stored on
each SkuAllocation while it worked (engine/allocation.py, engine/reconciliation.py
-- write-only, never read back into a calculation). Nothing is re-derived, so the
sentence cannot drift from the numbers, and the parts always add up:

    produced = W1A + carry-in spread + Run 1 + Run 2 + reconciliation adjustment

"Difference (T)" is the real conservation residual for the row:
    (produced + carry-out) - (FIN + carry-in)
Every quantity is rounded to 0.1 T at each step, so differences up to about
0.2 T are rounding; the run's health check flags anything above 0.5 T.

Figures in the sentence are rounded for reading. Where a sentence multiplies two
figures, enough decimals (up to 3) are shown that the displayed factors multiply
to the displayed product to within the display precision (0.05 T) -- a planner
checking it with a calculator must get the same answer, or the tool looks
wrong when it is not. Where that would take more than 3 decimals (a product
sitting near a rounding boundary -- about 1 in 8 Case C/D rows on real data)
the sentence says "~" instead of "=", which is what it is.

Contract:
    consumes: an openpyxl Workbook, the run's reconciled SkuAllocation rows
    produces: a new "Calculation Trace" worksheet in that workbook
"""
from __future__ import annotations

import math

from openpyxl.styles import Alignment, Font
from openpyxl.workbook.workbook import Workbook

from output import styling

SHEET_NAME = "Calculation Trace"

_COLUMNS = [
    ("Plant", 14), ("Line", 12), ("Linkcode", 12), ("Month", 9), ("Case", 8),
    ("Total FIN (T)", 13), ("Carry-in (T)", 12), ("Throughput (T/day)", 13), ("GE %", 8),
    ("Daily demand (T/day)", 13), ("Produced (T)", 12), ("Carry-out (T)", 12),
    ("Difference (T)", 12), ("How it was calculated", 110),
]
_EXPLANATION_COL = len(_COLUMNS)
_CHARS_PER_LINE = 100  # conservative: leaves room for proportional-font variation
_LINE_HEIGHT = 15

_FOOTNOTE = (
    "Difference = (produced + carry-out) - (FIN + carry-in). Quantities are rounded to 0.1 T at each step, "
    "so differences up to about 0.2 T are rounding; the run's health check flags anything above 0.5 T. "
    "Figures in the sentence are rounded for reading; the quantity columns are shown at the engine's 0.1 T resolution."
)


_DISPLAY_PRECISION_T = 0.0501  # half a display unit (0.1 T), plus float slack


_MAX_DECIMALS = 3


def _factor_decimals(a: float, b: float, product: float) -> tuple[int, str]:
    """(decimals, connector): fewest decimals (up to 3) at which the displayed
    factors multiply to the displayed product (shown to 0.1 T) to within the
    display precision, with "=". If even 3 decimals cannot, 3 decimals and "~"
    -- honest about the rounding rather than a long run of digits. (On the
    client's real data a fixed one-decimal display failed a calculator check on
    57 of 102 Case C/D rows.)"""
    for decimals in range(1, _MAX_DECIMALS + 1):
        if abs(round(a, decimals) * round(b, decimals) - round(product, 1)) <= _DISPLAY_PRECISION_T:
            return decimals, "="
    return _MAX_DECIMALS, "~"


def explain(a) -> str:
    """One concise explanation for one SkuAllocation row (see module docstring)."""
    parts: list[str] = []

    if a.carryover_fin_in > 0.05:
        text = f"Carry-in {a.carryover_fin_in:.1f} T: {a.wk1a:.1f} T in W1A"
        if a.carry_spread > 0.05:
            text += f", {a.carry_spread:.1f} T spread over W1-W4"
        unplaced = a.carryover_fin_in - a.wk1a - a.carry_spread
        if unplaced > 0.05:
            text += f", {unplaced:.1f} T could not be placed (no capacity)"
        if a.changeover_first:
            text += "; served first (month-end changeover rule)"
        parts.append(text + ".")

    fin, case = a.current_fin, a.moq_case
    if case == "No MOQ":
        parts.append("No MOQ: all FIN goes through Run 2.")
    elif case:
        if case == "A":
            text = f"Case A: FIN {fin:.1f} T < 1.5 x MOQ batch {a.moq_qty:.1f} T -> whole FIN in Run 1"
        elif case == "B":
            if a.moq_qty and fin >= 1.5 * a.moq_qty:
                text = (f"Case B: no DOS gap, FIN {fin:.1f} T >= 1.5 x MOQ batch {a.moq_qty:.1f} T "
                        f"-> half ({a.run1_target:.1f} T) in Run 1")
            else:
                text = f"Case B: no DOS gap -> one MOQ batch ({a.run1_target:.1f} T) in Run 1"
        else:  # C or D: sized from the DOS gap
            k, eq = _factor_decimals(a.dos_gap_days, a.daily_demand, a.dos_gap_qty)
            gap = f"DOS gap {a.dos_gap_days:.{k}f} d x {a.daily_demand:.{k}f} T/d {eq} {a.dos_gap_qty:.1f} T"
            if case == "C":
                text = f"Case C: {gap} < MOQ batch {a.moq_qty:.1f} T -> one MOQ batch ({a.run1_target:.1f} T) in Run 1"
            else:
                text = f"Case D: {gap} >= MOQ batch {a.moq_qty:.1f} T -> {a.run1_target:.1f} T in Run 1"
        if a.run1_placed < a.run1_target - 0.05:
            text += f", only {a.run1_placed:.1f} T fitted"
        parts.append(text + ".")

    if a.run2_placed > 0.05:
        parts.append(f"Run 2 placed {a.run2_placed:.1f} T.")
    if a.run2_deferred > 0.05:
        reason = "below the MOQ run-length floor" if a.run2_moq_blocked else "no capacity left"
        parts.append(f"{a.run2_deferred:.1f} T could not be placed in Run 2 ({reason}).")
    if abs(a.recon_adjustment) >= 0.05:
        parts.append(f"Reconciliation adjusted {a.recon_adjustment:+.1f} T.")
    parts.append(f"Carry-out {a.carryover_next:.1f} T.")
    return " ".join(parts)


def difference_t(a) -> float:
    """(produced + carry-out) - (FIN + carry-in): the row's real conservation residual."""
    return round((a.total_all + a.carryover_next) - (a.current_fin + a.carryover_fin_in), 2)


def write(workbook: Workbook, rows: list) -> None:
    ws = workbook.create_sheet(SHEET_NAME)  # appended last: existing sheet positions do not move

    for col, (title, width) in enumerate(_COLUMNS, start=1):
        ws.cell(row=1, column=col, value=title)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    styling.style_header_row(ws, ncols=len(_COLUMNS))

    ordered = sorted(rows, key=lambda a: (str(a.plant), str(a.line), str(a.link_code), a.period))
    for r, a in enumerate(ordered, start=2):
        text = explain(a)
        values = [
            a.plant, a.line, a.link_code, a.month_key, a.moq_case,
            round(a.current_fin, 1), round(a.carryover_fin_in, 1), a.throughput_per_day, a.ge_pct,
            round(a.daily_demand, 2), round(a.total_all, 1), round(a.carryover_next, 1),
            difference_t(a), text,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)
        ws.cell(row=r, column=9).number_format = "0.0%"
        ws.cell(row=r, column=_EXPLANATION_COL).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = _LINE_HEIGHT * max(1, math.ceil(len(text) / _CHARS_PER_LINE))

    last = len(ordered) + 1
    ws.freeze_panes = "E2"                                   # header row + Plant/Line/Linkcode/Month stay visible
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=_EXPLANATION_COL).column_letter}{last}"
    note = ws.cell(row=last + 2, column=1, value=_FOOTNOTE)
    note.font = Font(italic=True)
